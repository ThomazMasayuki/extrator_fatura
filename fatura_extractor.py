"""Extração de lançamentos de faturas de cartão em PDF (núcleo, sem interface).

Uso direto pela linha de comando:
    python extrator/fatura_extractor.py fatura.pdf [outra.pdf ...] --conta itau --saida dados/faturas
    python extrator/fatura_extractor.py fatura.pdf --senha 12345        # PDF protegido
    python extrator/fatura_extractor.py fatura.pdf --remover-senha --senha 12345   # só gera cópia sem senha
    python extrator/fatura_extractor.py fatura.pdf --texto               # texto bruto p/ diagnóstico

A interface gráfica (app_fatura.py) usa esta mesma classe.
"""
from __future__ import annotations

import argparse
import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

import pandas as pd
import pdfplumber
from pdfminer.pdfdocument import PDFPasswordIncorrect

try:  # pdfplumber >= 0.11 embrulha os erros do pdfminer
    from pdfplumber.utils.exceptions import PdfminerException
except ImportError:  # versões antigas
    PdfminerException = PDFPasswordIncorrect

MESES = {"JAN": 1, "FEV": 2, "MAR": 3, "ABR": 4, "MAI": 5, "JUN": 6,
         "JUL": 7, "AGO": 8, "SET": 9, "OUT": 10, "NOV": 11, "DEZ": 12}

# Linha de lançamento genérica:
#   DATA        27/08 | 27/08/2026 | 27.08.26 | 27 AGO
#   DESCRIÇÃO   qualquer texto
#   PAÍS        opcional (BB: "BR")
#   VALOR       1.234,56 com "R$" opcional; sinal '-' antes ou depois; ou sufixo C/D
RE_LINHA = re.compile(
    r"^(?:(?P<data>\d{2}[/.]\d{2}(?:[/.]\d{2,4})?)|(?P<dia_txt>\d{2})\s+(?P<mes_txt>"
    r"JAN|FEV|MAR|ABR|MAI|JUN|JUL|AGO|SET|OUT|NOV|DEZ)[A-Z]*)\s+"
    r"(?P<desc>.+?)\s+"
    r"(?:(?P<pais>[A-Z]{2}|10)\s+)?"
    r"(?P<s1>-\s?)?(?:R\$|US\$)?\s*(?P<s2>-\s?)?"
    r"(?P<valor>\d{1,3}(?:\.\d{3})*,\d{2})"
    r"(?P<s3>\s?-|\s[CD])?$",
    re.I)
RE_PARC = re.compile(r"(?:\bPARC(?:ELA)?\.?\s*)?\b(?P<n>\d{1,2})\s*(?:/|\sDE\s)\s*(?P<t>\d{1,2})\s*$", re.I)
RE_CATEGORIA = re.compile(r"^([A-ZÀ-Ú]{3,})\s*\.\s*(.*)$")
RE_MOEDA = re.compile(r"\b(USD|EUR|GBP|US\$)\b")
RE_PAGAMENTO = re.compile(r"PAGAMENTO|PAGTO|PGTO|PAG\s+FATURA|PAG\s+BOLETO|PAGAMENTO RECEBIDO", re.I)
RE_IGNORAR = re.compile(r"SALDO ANTERIOR|TOTAL\b|SUBTOTAL|LIMITE", re.I)
RE_FUTURO = re.compile(r"pr[oó]xim[ao]s? faturas?|parcelas? futuras?|lan[cç]amentos futuros|"
                       r"compras parceladas.{0,20}pr[oó]xim|a vencer", re.I)
RE_LISTA = re.compile(r"lan[cç]amentos|descri[cç][aã]o", re.I)
RE_GRUPO = re.compile(r"^[A-Za-zÀ-ú][A-Za-zÀ-ú &/,.-]{2,40}$")
RE_NAO_GRUPO = re.compile(r"cart[aã]o|fatura|p[aá]gina|total|data|valor|pa[ií]s|lan[cç]amento|"
                          r"resumo|limite|vencimento|encargos|juros", re.I)


# Rótulos do valor total, do mais específico ao mais genérico (sem acento, minúsculo)
ROTULOS_TOTAL = [r.split() for r in [
    "total desta fatura", "total da fatura", "valor desta fatura", "valor da fatura",
    "total da sua fatura", "o total da sua fatura", "valor total da fatura", "saldo desta fatura",
    "total a pagar", "valor a pagar", "valor total", "fatura atual",
]]
RE_VALOR_PALAVRA = re.compile(r"^-?(?:R\$)?-?(\d{1,3}(?:\.\d{3})*,\d{2})-?$")
RE_DATA_PALAVRA = re.compile(r"^(\d{2}[/.]\d{2}[/.]\d{2,4})$")


def _sem_acento(t: str) -> str:
    import unicodedata
    return unicodedata.normalize("NFKD", t).encode("ascii", "ignore").decode().lower().strip(":;.,")


def _data_br(v: str) -> date:
    d, m, a = map(int, re.split(r"[/.]", v))
    return date(a + 2000 if a < 100 else a, m, d)


def _valor_perto(paginas: list[list[dict]], rotulos: list[list[str]], padrao: re.Pattern, conv):
    """Acha um rótulo pela posição das palavras e devolve o valor mais próximo
    à direita (mesma linha) ou logo abaixo (mesma coluna). Independe do banco."""
    for palavras in paginas:
        norm = [_sem_acento(w["text"]) for w in palavras]
        for rotulo in rotulos:
            n = len(rotulo)
            for i in range(len(palavras) - n + 1):
                if norm[i:i + n] != rotulo:
                    continue
                ini, fim = palavras[i], palavras[i + n - 1]
                if abs(fim["top"] - ini["top"]) > 4:   # rótulo quebrado em duas linhas: ignora
                    continue
                seguinte = norm[i + n] if i + n < len(norm) else ""
                if seguinte in ("anterior", "passada", "minima", "minimo"):   # "total da fatura anterior"
                    continue
                x0, x1, topo, base = ini["x0"], fim["x1"], ini["top"], fim["bottom"]
                melhor = None
                for w in palavras:
                    m = padrao.match(w["text"].replace(" ", ""))
                    if not m:
                        continue
                    if abs(w["top"] - topo) <= 4 and w["x0"] >= x1 - 1:            # à direita
                        dist = w["x0"] - x1
                    elif 0 <= w["top"] - base <= 60 and w["x1"] >= x0 - 40 and w["x0"] <= x1 + 40:  # abaixo
                        dist = (w["top"] - base) * 2 + abs(w["x0"] - x0) * 0.2
                    else:
                        continue
                    if melhor is None or dist < melhor[0]:
                        melhor = (dist, m.group(1))
                if melhor:
                    try:
                        return conv(melhor[1])
                    except ValueError:
                        continue
    return None


def _brl(v: float) -> str:
    return "R$ " + f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


class SenhaNecessaria(Exception):
    """O PDF é protegido e a senha não foi informada ou está errada."""


def abrir_pdf(caminho: str | Path, senhas: list[str] | tuple[str, ...] = ()) -> tuple[pdfplumber.PDF, str]:
    """Tenta abrir sem senha e depois com cada senha informada. Devolve (pdf, senha_que_funcionou)."""
    for senha in ("", *senhas):
        try:
            return pdfplumber.open(caminho, password=senha), senha
        except (PDFPasswordIncorrect, PdfminerException) as e:
            causa = e.__cause__ or (e.args[0] if e.args else None)
            if isinstance(e, PDFPasswordIncorrect) or isinstance(causa, PDFPasswordIncorrect):
                continue
            raise
    raise SenhaNecessaria(f"{Path(caminho).name} é protegido por senha.")


def precisa_senha(caminho: str | Path, senhas: list[str] | tuple[str, ...] = ()) -> bool:
    try:
        pdf, _ = abrir_pdf(caminho, senhas)
        pdf.close()
        return False
    except SenhaNecessaria:
        return True


def remover_senha(caminho: str | Path, senha: str, destino: str | Path | None = None) -> Path:
    """Salva uma cópia do PDF sem senha (o original não é alterado)."""
    from pypdf import PdfReader, PdfWriter

    caminho = Path(caminho)
    destino = Path(destino) if destino else caminho.with_name(f"{caminho.stem}_sem_senha.pdf")
    leitor = PdfReader(caminho)
    if leitor.is_encrypted and not leitor.decrypt(senha):
        raise SenhaNecessaria("Senha incorreta.")
    escritor = PdfWriter(clone_from=leitor)
    with open(destino, "wb") as f:
        escritor.write(f)
    return destino


def exportar_texto(caminho: str | Path, senha: str | None = None, destino: str | Path | None = None) -> Path:
    """Salva o texto bruto que o leitor enxerga (para diagnosticar layouts novos)."""
    caminho = Path(caminho)
    destino = Path(destino) if destino else caminho.with_suffix(".txt")
    pdf, _ = abrir_pdf(caminho, [senha] if senha else [])
    with pdf:
        paginas = [f"===== página {i} =====\n{p.extract_text() or ''}" for i, p in enumerate(pdf.pages, 1)]
    destino.write_text("\n".join(paginas), encoding="utf-8")
    return destino


def _num(txt: str) -> float:
    return float(txt.replace(" ", "").replace(".", "").replace(",", "."))


def _somar_meses(d: date, meses: int) -> date:
    m = d.month - 1 + meses
    ano, mes = d.year + m // 12, m % 12 + 1
    dias = [31, 29 if ano % 4 == 0 and (ano % 100 or ano % 400 == 0) else 28,
            31, 30, 31, 30, 31, 31, 30, 31, 30, 31][mes - 1]
    return date(ano, mes, min(d.day, dias))


@dataclass
class Fatura:
    arquivo: Path
    titular: str | None = None
    numero_cartao: str | None = None
    valor_total: float | None = None
    vencimento: date | None = None
    pagamento_minimo: float | None = None
    pagamento_parcelado: float | None = None
    limite_total: float | str | None = None
    limite_utilizado: float | None = None
    pagamentos: float = 0.0
    lancamentos: pd.DataFrame = field(default_factory=pd.DataFrame)
    avisos: list[str] = field(default_factory=list)

    @property
    def total_extraido(self) -> float:
        return float(self.lancamentos["Valor"].sum()) if not self.lancamentos.empty else 0.0

    @property
    def diferenca(self) -> float | None:
        return None if self.valor_total is None else round(self.valor_total - self.total_extraido, 2)


class FaturaExtractor:
    def __init__(self, pdf_path: str | Path, remover_duplicados: bool = False, senha: str | None = None):
        self.pdf_path = Path(pdf_path)
        self.remover_duplicados = remover_duplicados
        self.senha = senha
        self._palavras: list[list[dict]] = []
        self.fatura = Fatura(arquivo=self.pdf_path)

    # ------------------------------------------------------------------ texto
    def extract_text(self) -> str:
        partes = []
        pdf, _ = abrir_pdf(self.pdf_path, [self.senha] if self.senha else [])
        with pdf:
            for i, page in enumerate(pdf.pages):
                partes.append(page.extract_text() or "")  # página só com imagem devolve None
                if i < 2:  # cabeçalho/resumo fica nas primeiras páginas
                    self._palavras.append(page.extract_words(keep_blank_chars=False))
        texto = "\n".join(partes)
        if not texto.strip():
            raise ValueError("O PDF não tem texto extraível (provavelmente é imagem escaneada).")
        return texto

    # ------------------------------------------------------------ cabeçalho
    def extract_basic_info(self, texto: str) -> None:
        """Campos do cabeçalho. Cada campo tem várias redações possíveis (uma por banco)."""
        f = self.fatura

        def achar(padroes, flags=re.I):
            for p in padroes:
                if (m := re.search(p, texto, flags)):
                    return m.group(1).strip()
            return None

        VAL = r"(?:R\$\s*)?(\d{1,3}(?:\.\d{3})*,\d{2})"
        f.titular = achar([r"Titular\s+([A-ZÀ-Ú][A-ZÀ-Ú ]+?)\s*\n",
                           r"\d{2}\s*-\s*([A-ZÀ-Ú][A-ZÀ-Ú ]+?)\s+Cart[aã]o\s+N",
                           r"([A-ZÀ-Ú ]{6,})\s+Resumo da fatura em R\$"], flags=0)
        f.numero_cartao = achar([r"Cart[aã]o\s+(?:N\.?\s*)?(?:final\s*)?([0-9X\.\s]{4,}?)\s*[\n)]"])
        if (v := achar([rf"(?s)O total da sua fatura é:.{{0,40}}?{VAL}", rf"Total (?:desta|da) fatura(?!\s*(?:anterior|passada))[^\d\n]{{0,30}}{VAL}",
                        rf"Valor total(?: da fatura)?[^\d\n]{{0,30}}{VAL}", rf"Total a pagar[^\d\n]{{0,30}}{VAL}"])):
            f.valor_total = _num(v)
        if (v := achar([r"Vencimento[^\d]{0,40}(\d{2}[/.]\d{2}[/.]\d{2,4})"])):
            d, m, a = map(int, re.split(r"[/.]", v))
            f.vencimento = date(a + 2000 if a < 100 else a, m, d)
        elif (v := achar([r"Vencimento[^\d]{0,40}(\d{2}\s+[A-Za-z]{3}\s+\d{4})"])):   # Nubank: 15 SET 2026
            d, m, a = v.split()
            if m.upper()[:3] in MESES:
                f.vencimento = date(int(a), MESES[m.upper()[:3]], int(d))
        # Fallback por posição: rótulo e valor em caixas separadas (ex.: BB, 1ª página)
        if f.valor_total is None:
            f.valor_total = _valor_perto(self._palavras, ROTULOS_TOTAL, RE_VALOR_PALAVRA, _num)
        if f.vencimento is None:
            f.vencimento = _valor_perto(self._palavras, [["vencimento"]], RE_DATA_PALAVRA, _data_br)

        if (v := achar([rf"Pagamento mínimo[^\d]{{0,30}}{VAL}"])):
            f.pagamento_minimo = _num(v)
        if (v := achar([rf"Parcelas fixas[^\d]{{0,30}}{VAL}"])):
            f.pagamento_parcelado = _num(v)
        if (v := achar([r"Limite (?:total )?utilizado\s+(?:R\$\s*)?([0-9\.,]+)"])):
            f.limite_utilizado = _num(v)
        if (v := achar([r"Limite total(?: de crédito)?\s+(?:R\$\s*)?([0-9\.]+,\d{2}|Flexível)"])):
            f.limite_total = v if v == "Flexível" else _num(v)

        if f.vencimento is None:
            f.avisos.append("Vencimento não encontrado: o ano das compras foi estimado pela data de hoje.")

    # ----------------------------------------------------------------- datas
    def _datas(self, dia: int, mes: int, ano: int | None, parcela: str | None) -> tuple[date, date]:
        """Resolve o ano da compra e a data em que a parcela cai.

        A compra original é datada no passado; a parcela N/M é cobrada N-1 meses depois.
        Sem ano explícito, usa o mais recente em que a parcela não cai depois do vencimento.
        """
        ref = self.fatura.vencimento or date.today()
        n = int(parcela.split("/")[0]) if parcela else 1
        anos = [ano] if ano else [ref.year, ref.year - 1, ref.year - 2, ref.year - 3]
        for a in anos:
            try:
                compra = date(a, mes, dia)
            except ValueError:  # 29/02 em ano não bissexto
                continue
            cobranca = _somar_meses(compra, n - 1)
            if ano or cobranca <= ref:
                return compra, cobranca
        compra = date(anos[0], mes, min(dia, 28))
        return compra, _somar_meses(compra, n - 1)

    # ------------------------------------------------------------ lançamentos
    def extract_purchases(self, texto: str) -> None:
        """Leitor genérico: qualquer linha 'DATA  DESCRIÇÃO  VALOR' é um lançamento.

        Não depende de títulos de seção, então serve para layouts de bancos diferentes.
        """
        linhas = [" ".join(l.split()) for l in texto.split("\n")]
        itens, vistos, duplicados, ignoradas = [], set(), 0, 0
        pagamentos = 0.0
        grupo = None          # linha de grupo acima das compras (BB: "Restaurantes")
        futuro = False        # dentro de quadro de parcelas futuras / próximas faturas

        for i, linha in enumerate(linhas):
            if not linha:
                continue
            if RE_FUTURO.search(linha):
                futuro = True
                continue
            if RE_LISTA.search(linha):
                futuro = False
                continue

            m = RE_LINHA.match(linha)
            if not m:
                # linha de grupo (BB: "Restaurantes"): curta, sem dígitos e seguida de um lançamento
                prox_util = next((l for l in linhas[i + 1:] if l), "")
                if (RE_GRUPO.match(linha) and not RE_NAO_GRUPO.search(linha)
                        and not RE_CATEGORIA.match(linha) and RE_LINHA.match(prox_util)):
                    grupo = linha.strip().capitalize()
                continue
            if futuro:
                ignoradas += 1
                continue

            g = m.groupdict()
            desc = g["desc"].strip()
            # descrição precisa ter palavra de verdade (evita linhas de resumo "10/09/2026 R$ 137,25 R$ 20,59")
            if (not re.search(r"[A-Za-zÀ-ú]{2,}", re.sub(r"R\$|US\$", "", desc))
                    or re.search(r"R\$\s*\d", desc) or RE_IGNORAR.search(desc)):
                continue

            # sinal: '-' antes/depois do valor ou sufixo C (crédito) = entrada no cartão
            credito = bool(g["s1"] or g["s2"] or (g["s3"] and g["s3"].strip() in ("-", "C")))
            valor = _num(g["valor"]) * (-1 if credito else 1)

            if credito and RE_PAGAMENTO.search(desc):
                pagamentos += -valor   # pagamento da fatura anterior: já entra pelo extrato
                continue

            # parcela no fim da descrição: "LOJA 02/10", "LOJA PARC 02/10", "LOJA PARCELA 2 DE 10"
            parcela = None
            if (pm := RE_PARC.search(desc)):
                n, t = int(pm.group("n")), int(pm.group("t"))
                if 1 <= n <= t <= 48 and t > 1:
                    parcela = f"{n:02d}/{t:02d}"
                    desc = desc[:pm.start()].strip(" -")

            # moeda estrangeira: país diferente de BR ou linha seguinte com USD/EUR
            prox = linhas[i + 1] if i + 1 < len(linhas) else ""
            pais = g.get("pais")
            internacional = (pais and pais not in ("BR", "10")) or bool(RE_MOEDA.search(prox))

            categoria = grupo
            if (c := RE_CATEGORIA.match(prox)) and not RE_LINHA.match(prox):
                categoria = c.group(1).capitalize()      # Itaú: "RESTAURANTE . CIDADE"

            dia, mes, ano = self._dia_mes_ano(g)
            chave = (dia, mes, desc, parcela, valor)
            if self.remover_duplicados and chave in vistos:
                duplicados += 1
                continue
            vistos.add(chave)
            compra, cobranca = self._datas(dia, mes, ano, parcela)
            itens.append({
                "Data": cobranca,
                "Descricao": desc,
                "Valor": valor,
                "Parcela": parcela or "",
                "Categoria_banco": categoria or ("Internacional" if internacional else ""),
                "Tipo": "Internacional" if internacional else "Nacional",
                "Data_compra": compra,
            })

        df = pd.DataFrame(itens, columns=["Data", "Descricao", "Valor", "Parcela",
                                          "Categoria_banco", "Tipo", "Data_compra"])
        if not df.empty:
            df["Data"] = pd.to_datetime(df["Data"])
            df["Data_compra"] = pd.to_datetime(df["Data_compra"])
            df = df.sort_values(["Data", "Descricao"]).reset_index(drop=True)
        self.fatura.lancamentos = df
        self.fatura.pagamentos = round(pagamentos, 2)

        if df.empty:
            self.fatura.avisos.append(
                "Nenhum lançamento encontrado. Use 'Exportar texto' e envie o .txt para ajustar o leitor.")
        if pagamentos:
            self.fatura.avisos.append(
                f"Pagamento(s) da fatura anterior ({_brl(pagamentos)}) não entraram como lançamento "
                "— eles já aparecem no extrato bancário.")
        if ignoradas:
            self.fatura.avisos.append(f"{ignoradas} linha(s) de parcelas futuras/próximas faturas foram ignoradas.")
        if duplicados:
            self.fatura.avisos.append(f"{duplicados} linha(s) repetida(s) idênticas foram ignoradas.")
        if (dif := self.fatura.diferenca) is not None and abs(dif) > 0.01:
            self.fatura.avisos.append(
                f"Soma extraída difere do total da fatura em {_brl(dif)}. Pode ser saldo anterior, "
                "juros/IOF/anuidade sem data ou linha não reconhecida.")

    @staticmethod
    def _dia_mes_ano(g: dict) -> tuple[int, int, int | None]:
        if g.get("mes_txt"):
            return int(g["dia_txt"]), MESES[g["mes_txt"].upper()[:3]], None
        partes = [int(x) for x in re.split(r"[/.]", g["data"])]
        ano = None
        if len(partes) == 3:
            ano = partes[2] + 2000 if partes[2] < 100 else partes[2]
        return partes[0], partes[1], ano

    # ---------------------------------------------------------------- fluxo
    def process(self) -> Fatura:
        texto = self.extract_text()
        self.extract_basic_info(texto)
        self.extract_purchases(texto)
        return self.fatura


# =============================================================================
# Exportação (formato lido pelo planejador: 1ª aba = lançamentos)
# =============================================================================
def nome_saida(fatura: Fatura, conta: str) -> str:
    ref = fatura.vencimento.strftime("%Y-%m") if fatura.vencimento else fatura.arquivo.stem
    conta = re.sub(r"[^\w\-]+", "_", conta.strip().lower()) or "cartao"
    return f"{conta}__{ref}.xlsx"


def exportar_excel(fatura: Fatura, caminho: str | Path) -> Path:
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    caminho = Path(caminho)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    df = fatura.lancamentos.copy()

    resumo = pd.DataFrame({
        "Informação": ["Arquivo", "Titular", "Cartão", "Vencimento", "Valor total da fatura",
                       "Soma dos lançamentos extraídos", "Diferença", "Pagamentos recebidos (fora)",
                       "Pagamento mínimo",
                       "Parcelas fixas", "Limite total", "Limite utilizado", "Qtd. lançamentos"],
        "Valor": [fatura.arquivo.name, fatura.titular, fatura.numero_cartao,
                  fatura.vencimento, fatura.valor_total, round(fatura.total_extraido, 2),
                  fatura.diferenca, fatura.pagamentos, fatura.pagamento_minimo, fatura.pagamento_parcelado,
                  fatura.limite_total, fatura.limite_utilizado, len(df)],
    })
    por_cat = (df.assign(Categoria_banco=df["Categoria_banco"].replace("", "Sem categoria"))
                 .groupby("Categoria_banco", as_index=False)
                 .agg(Total=("Valor", "sum"), Qtd=("Valor", "count"))
                 .sort_values("Total", ascending=False)) if not df.empty else pd.DataFrame()

    with pd.ExcelWriter(caminho, engine="openpyxl") as w:
        df.to_excel(w, sheet_name="Lancamentos", index=False)
        resumo.to_excel(w, sheet_name="Resumo", index=False)
        if not por_cat.empty:
            por_cat.to_excel(w, sheet_name="Por Categoria", index=False)
        if fatura.avisos:
            pd.DataFrame({"Aviso": fatura.avisos}).to_excel(w, sheet_name="Avisos", index=False)

        for ws in w.book.worksheets:
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="1F4E79")
            for col in ws.columns:
                largura = max(len(str(c.value or "")) for c in col)
                ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(12, largura + 2), 60)
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                        c.number_format = "#,##0.00" if isinstance(c.value, float) else "0"
                    elif hasattr(c.value, "year"):
                        c.number_format = "DD/MM/YYYY"
    return caminho


def main() -> None:
    ap = argparse.ArgumentParser(description="Extrai lançamentos de faturas em PDF para Excel.")
    ap.add_argument("pdfs", nargs="+")
    ap.add_argument("--conta", default="cartao", help="nome da conta (vira prefixo do arquivo)")
    ap.add_argument("--saida", default=".", help="pasta de saída")
    ap.add_argument("--senha", help="senha do PDF (se protegido); sem ela, será perguntada")
    ap.add_argument("--remover-senha", action="store_true", help="apenas salvar cópia sem senha")
    ap.add_argument("--texto", action="store_true", help="apenas salvar o texto bruto (.txt) para diagnóstico")
    args = ap.parse_args()
    for pdf in args.pdfs:
        senha = args.senha
        if precisa_senha(pdf, [senha] if senha else []):
            import getpass
            senha = getpass.getpass(f"Senha de {Path(pdf).name}: ")
        if args.texto:
            print(f"{Path(pdf).name}: texto -> {exportar_texto(pdf, senha)}")
            continue
        if args.remover_senha:
            print(f"{Path(pdf).name}: cópia sem senha -> {remover_senha(pdf, senha or '')}")
            continue
        fat = FaturaExtractor(pdf, senha=senha).process()
        destino = exportar_excel(fat, Path(args.saida) / nome_saida(fat, args.conta))
        print(f"{Path(pdf).name}: {len(fat.lancamentos)} lançamentos -> {destino}")
        for a in fat.avisos:
            print(f"  ! {a}")


if __name__ == "__main__":
    main()
